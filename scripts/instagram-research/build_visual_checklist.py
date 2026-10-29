#!/usr/bin/env python3
"""Build the worklist for checking the still-ambiguous rows by eye on Instagram.

Every row here is one this project could not settle from outside a login: the
candidate accounts found so far, the address the Maps audit list records, what
exactly to look for, and blank columns for the verdict. Nothing is pre-filled as
a result - the point of the sheet is that a person decides.
"""
import json, re
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
SRC = Path("/Users/hashiramototetsuya/仕事/Coding仕様書/tokyo_guide_maps_link_review_unified_557_instagram.xlsx")
OUT = SRC.with_name("tokyo_guide_instagram_目視確認リスト.xlsx")

IG = "https://www.instagram.com/"
HANDLE = re.compile(r"^[A-Za-z0-9._]{1,30}$")

CLASS = {
    "A": "A 候補あり・プロフィールに所在地が無い",
    "B": "B 候補が複数あり現行が不明",
    "C": "C 店主・店長の個人アカウントのみ",
    "D": "D 当該支店のアカウントが無い",
    "E": "E アカウントの所在地が台帳と不一致",
    "F": "F 施設自身のアカウントが無い",
    "G": "G 候補アカウントが見つかっていない",
}
ORDER = ["A", "B", "E", "D", "C", "F", "G"]
TODO = {
    "A": ("候補アカウントを開き、プロフィール・ハイライト・固定投稿・直近の投稿のどこかに、"
          "右の「住所参考」と一致する住所・最寄駅・電話番号・区名が出ているかを見る。"
          "出ていれば公式と確認。別の場所が書かれていればNG。"),
    "B": ("候補が複数あるので、各アカウントの最終投稿日・プロフィールのリンク先・フォロワー数を見比べ、"
          "現在運用されている方を特定する。両方生きている場合は、店名表記が正式名と一致する方を優先。"),
    "C": ("店の公式アカウントが別に存在しないかを先に確認する。無い場合、この店主・店長アカウントを"
          "公式として扱うかどうかは方針判断（現在は不採用）。採用するなら判定欄にその旨を記入。"),
    "D": ("ブランド側アカウントの投稿・ハイライト・プロフィールのリンク先に、この支店の告知や"
          "店舗一覧があるかを見る。支店が明記されていればブランド共通アカウントとして採用可。"),
    "E": ("アカウントが示す所在地と台帳の住所が食い違っている。移転したのか、別店舗なのか、"
          "台帳側の住所が古いのかを投稿履歴で確認する。"),
    "F": ("施設自身のアカウントが本当に存在しないかの最終確認。無いと確認できれば"
          "「アカウント無しと確認」を選び、NOT_FOUNDとして確定する。"),
    "G": ("まず店名での検索をやり直す（表記ゆれ・英字表記・旧店名）。見つかったら"
          "プロフィールの所在地が台帳の住所と一致するかを見る。"),
}
OVERRIDE = {
    "L120": "F", "L136": "F", "L139": "F", "L142": "F", "L457": "F", "L376": "F", "L547": "F",
    "L151": "E", "L357": "D", "L347": "G", "L475": "E",
}
# accounts this project did find during the research but which are named only in the
# prose of a row's reason; listed here so the sheet can link them for the eye check
EXTRA_CANDS = {
    "L022": ["tokyo.tonkatsu.gavuu"],
    "L087": ["amam_dacotan_omotesando", "amam_dacotan_omotedando", "dacotan.cafe.bake"],
    "L097": ["niihao_kamata"],
    "L120": ["yasukuni.official", "yasukunijinja"],
    "L136": ["meijijingu.wedding"],
    "L139": ["yasukuni.official"],
    "L151": ["akaiheya16"],
    "L322": ["santora_ramen", "sugaraaa_santora"],
    "L346": ["shogunburger", "shogunburger_hokkaido"],
    "L357": ["hama_ginza"],
    "L376": ["asakusapelicancafe"],
    "L475": ["chedihanshin"],
    "L495": ["kitatougarashi"],
    "L547": ["hakuba_kurashita_carstay", "hakuba_kurashita_camp"],
}
NOTE = {
    "L347": "この行のGoogleマップのピンが、台帳側の「ビフテキのカワムラ 銀座」とは別の場所を指している。"
            "銀座の「かわむら」は同名の別店（ステーキ店）の可能性があるため、まずどちらの店かを確定する。",
    "L087": "2アカウントとも住所は北青山3-7-6で一致している。どちらが現行かだけを見れば足りる。",
    "L298": "食べログ掲載の候補 1500_millecinquecento を含めて確認する。",
}


def main() -> int:
    ws = load_workbook(SRC, read_only=True)["557件統合監査"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    head = rows[0]
    body = [r for r in rows[1:] if any(c is not None for c in r)]
    I = {k: head.index(k) for k in ("掲載ID", "施設名（訂正提案反映）", "掲載地域", "エリア", "カテゴリ",
                                    "Instagram判定", "Instagram根拠", "Instagram根拠URL",
                                    "住所参考（過去の独立結果）", "原記録URL（監査証跡）")}
    weak = {}
    for r in json.loads((HERE / "stage6_results.json").read_text(encoding="utf-8")):
        weak[r["id"]] = sorted({w["handle"] for w in r.get("weak_candidates", [])})

    items = []
    for r in body:
        if not str(r[I["Instagram判定"]]).startswith("AMBIGUOUS"):
            continue
        lid = r[I["掲載ID"]]
        ev = str(r[I["Instagram根拠"]] or "")
        cands: list[str] = []
        url = str(r[I["Instagram根拠URL"]] or "")
        if url.startswith(IG):
            h = url[len(IG):].strip("/")
            if HANDLE.match(h):
                cands.append(h)
        for marker in ("Candidates:", "Printed:", "Seen:"):
            if marker in ev:
                tail = ev.split(marker, 1)[1].split(".")[0]
                for t in re.split(r"[,、\s]+", tail):
                    t = t.strip()
                    if HANDLE.match(t) and "." not in t[:1] and t not in cands and len(t) > 3:
                        cands.append(t)
        for h in weak.get(lid, []):
            if h not in cands:
                cands.append(h)
        for h in EXTRA_CANDS.get(lid, []):
            if h not in cands:
                cands.append(h)

        e = ev.lower()
        cls = OVERRIDE.get(lid)
        if cls is None:
            if "proprietor" in e or "manager's" in e or "personal account" in e:
                cls = "C"
            elif "two accounts" in e or "near-identical" in e:
                cls = "B"
            elif "rather than" in e and "places it" in e:
                cls = "E"
            elif "no account for the" in e and ("branch" in e or "main shop" in e):
                cls = "D"
            elif "no account of its own" in e:
                cls = "F"
            elif cands:
                cls = "A"
            else:
                cls = "G"
        items.append({
            "cls": cls, "id": lid, "name": r[I["施設名（訂正提案反映）"]],
            "region": r[I["掲載地域"]], "area": r[I["エリア"]], "cat": r[I["カテゴリ"]],
            "addr": r[I["住所参考（過去の独立結果）"]] or "", "cands": cands[:3],
            "why": ev, "maps": r[I["原記録URL（監査証跡）"]] or "",
        })
    items.sort(key=lambda x: (ORDER.index(x["cls"]), x["region"] != "Tokyo", x["area"] or "", x["id"]))

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "目視確認リスト"
    cols = ["作業No", "分類", "掲載ID", "施設名", "地域", "エリア", "カテゴリ", "住所参考（照合の基準）",
            "候補1", "候補2", "候補3", "確認すべき点", "【入力】判定", "【入力】確定アカウント",
            "【入力】確認できた根拠", "【入力】確認日", "これまでに調べたこと（保留理由）", "Google Maps URL"]
    widths = [7, 30, 9, 30, 9, 16, 14, 34, 34, 34, 34, 54, 20, 22, 40, 13, 80, 44]
    ws1.append(cols)
    for i, c in enumerate(ws1[1], start=1):
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="151513" if i < 13 or i > 16 else "1F6F43")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    thin = Side(style="thin", color="DDDDDD")
    for n, it in enumerate(items, start=1):
        cands = [f"{IG}{h}/" for h in it["cands"]] + ["", "", ""]
        ws1.append([n, CLASS[it["cls"]], it["id"], it["name"], it["region"], it["area"], it["cat"], it["addr"],
                    cands[0], cands[1], cands[2],
                    (NOTE[it["id"]] + " " if it["id"] in NOTE else "") + TODO[it["cls"]],
                    "", "", "", "", it["why"], it["maps"]])
        row = ws1.max_row
        for col in (9, 10, 11, 18):
            c = ws1.cell(row=row, column=col)
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value
                c.font = Font(color="0563C1", underline="single")
        for col in (12, 17):
            ws1.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(13, 17):
            ws1.cell(row=row, column=col).fill = PatternFill("solid", fgColor="EAF5EF")
        for col in range(1, len(cols) + 1):
            ws1.cell(row=row, column=col).border = Border(bottom=thin)
    for i, w in enumerate(widths, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = "D2"
    ws1.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws1.max_row}"
    dv = DataValidation(type="list", formula1='"公式と確認,別の店・非公式,アカウント無しと確認,判断できず"', allow_blank=True)
    dv.prompt = "見た結果を選んでください"
    dv.promptTitle = "判定"
    ws1.add_data_validation(dv)
    dv.add(f"M2:M{ws1.max_row}")

    ws2 = wb.create_sheet("使い方・判定基準")
    rows2 = [
        ("このリストは何か",
         "557件のうち、Instagramの公式アカウントを未ログインの調査では確定できなかった74件です。"
         "候補アカウントのURL・照合の基準になる住所・何を見れば決まるかを行ごとにまとめています。"),
        ("作業のしかた",
         "分類Aから順に並べてあります。Aは候補アカウントを開いて住所が出ているか見るだけなので最も速く、"
         "下に行くほど判断が重くなります。『候補1〜3』のURLをクリックしてアカウントを開き、"
         "『確認すべき点』に書いた観点で見て、緑色の4列（判定・確定アカウント・根拠・確認日）に記入してください。"),
        ("判定の選択肢",
         "公式と確認＝この施設の公式アカウントだと確認できた／別の店・非公式＝候補は別の店か非公式だった／"
         "アカウント無しと確認＝この施設にアカウントは無いと確認できた／判断できず＝見ても決まらなかった。"),
        ("「公式と確認」の条件",
         "次のどちらかが満たせれば確認済みとして台帳に反映できます。(1) 施設自身のウェブサイトのページから"
         "そのアカウントへリンクがある。(2) アカウント自身のプロフィール・ハイライト・固定投稿に施設名があり、"
         "かつ住所・最寄駅・電話番号・区名のいずれかが『住所参考』と一致している。"),
        ("採用していない根拠",
         "店名が同じだけ、フォロワーが多いだけ、第三者のまとめ記事に載っているだけ、写真が似ているだけ、"
         "の4つは根拠として採用していません。『根拠』欄には、プロフィールのどこに何と書いてあったかを"
         "一行で書いていただけると、そのまま台帳の根拠列に転記できます。"),
        ("確定アカウントの書き方",
         "候補1〜3のどれかなら、その番号またはアカウント名（@なしでも可）を書いてください。"
         "候補以外のアカウントが正しかった場合は、そのアカウント名を書いてください。"),
        ("分類の意味",
         "A 候補あり・プロフィールに所在地が無い／B 候補が複数あり現行が不明／"
         "E アカウントの所在地が台帳と不一致／D 当該支店のアカウントが無い／"
         "C 店主・店長の個人アカウントのみ／F 施設自身のアカウントが無い／G 候補が見つかっていない。"),
        ("方針の判断が要る行",
         "分類C（5件）は、店主・店長の個人アカウントしか存在しない行です。現在の台帳では"
         "「個人アカウントは店舗の公式として扱わない」方針で不採用にしています。"
         "採用する方針に変えるなら、5件まとめて切り替えるのが一貫します。"),
        ("戻していただいたあと",
         "記入後のファイルをそのまま渡していただければ、判定を台帳・サイトデータ・557件リストに反映し、"
         "テストと表示確認までやり直します。根拠欄の文面もそのまま根拠列に入れます。"),
        ("件数", f"{len(items)}件（作成日 2026-10-29）"),
    ]
    for r in rows2:
        ws2.append(list(r))
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 118
    for row in range(1, ws2.max_row + 1):
        ws2.cell(row=row, column=1).font = Font(bold=True)
        ws2.cell(row=row, column=1).alignment = Alignment(vertical="top")
        ws2.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT)
    from collections import Counter
    print("written:", OUT)
    print("rows:", len(items))
    for c, n in sorted(Counter(i["cls"] for i in items).items()):
        print(f"  {CLASS[c]}: {n}")
    print("rows with no candidate url:", sum(1 for i in items if not i["cands"]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
