#!/usr/bin/env python3
"""Write the final workbook: the master 526 rows with the Instagram columns added.

The original eight master columns and the row order are reproduced unchanged;
the Instagram columns are appended to the right.
"""
import csv, json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
MASTER = Path("/Users/hashiramototetsuya/Documents/Claude_新規事業/tokyo-guide/publish526_2026-10-22/deliverable/FINAL_APPLIED_526.csv")
OUT = HERE / "ENCOUNTER_JAPAN_526_INSTAGRAM.xlsx"

MASTER_COLS = ["live_order", "region", "name", "en", "desc", "area", "cat", "url"]
NEW_COLS = ["instagram_url", "instagram_username", "account_type", "basis", "status",
            "evidence", "evidence_url", "checked_at"]
WIDTHS = {"live_order": 10, "region": 12, "name": 34, "en": 26, "desc": 34, "area": 18,
          "cat": 16, "url": 44, "instagram_url": 38, "instagram_username": 24,
          "account_type": 13, "basis": 12, "status": 12, "evidence": 70, "evidence_url": 44, "checked_at": 12}
STATUS_FILL = {
    "VERIFIED": PatternFill("solid", fgColor="FFEADB"),
    "AMBIGUOUS": PatternFill("solid", fgColor="FFF6E0"),
    "NOT_FOUND": PatternFill("solid", fgColor="F1F0EC"),
}


def main() -> int:
    master = list(csv.DictReader(MASTER.open(encoding="utf-8-sig")))
    ledger = {r["live_order"]: r for r in json.loads((HERE / "ledger.json").read_text(encoding="utf-8"))}
    assert len(master) == 526 == len(ledger)

    wb = Workbook()
    ws = wb.active
    ws.title = "526_instagram"

    head = MASTER_COLS + NEW_COLS
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="151513")
        cell.alignment = Alignment(vertical="center")

    for m in master:
        rec = ledger[int(m["live_order"])]
        assert rec["name"] == m["name"] and rec["maps_url"] == m["url"], m["name"]
        ws.append([m[c] for c in MASTER_COLS] + [rec[c] for c in NEW_COLS])

    for i, col in enumerate(head, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(col, 18)
    status_col = head.index("status") + 1
    ig_col = head.index("instagram_url") + 1
    maps_col = head.index("url") + 1
    for row in range(2, ws.max_row + 1):
        st = ws.cell(row=row, column=status_col)
        st.fill = STATUS_FILL[st.value]
        st.font = Font(bold=True)
        for col in (ig_col, maps_col, head.index("evidence_url") + 1):
            c = ws.cell(row=row, column=col)
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value
                c.font = Font(color="0563C1", underline="single")
        ws.cell(row=row, column=head.index("evidence") + 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # second sheet: the counts, so the workbook explains itself
    s = wb.create_sheet("summary")
    from collections import Counter
    st = Counter(r["status"] for r in ledger.values())
    ac = Counter(r["account_type"] for r in ledger.values() if r["status"] == "VERIFIED")
    ba = Counter(r["basis"] for r in ledger.values() if r["status"] == "VERIFIED")
    rows = [
        ("Listings researched", 526),
        ("VERIFIED - official Instagram confirmed", st["VERIFIED"]),
        ("   confirmed by a link on the venue's own website", ba["own-site"]),
        ("   confirmed by the account's own profile (name + place)", ba["profile"]),
        ("   of which the venue's / branch's own account", ac["store"]),
        ("   of which a brand or operator account", ac["brand"]),
        ("AMBIGUOUS - candidate found, officiality not confirmable", st["AMBIGUOUS"]),
        ("NOT_FOUND - no candidate, reason recorded on the row", st["NOT_FOUND"]),
        ("", ""),
        ("Rule for VERIFIED", "either the account is linked from a page on the venue's own website, or the account's own profile names this venue and places it where this guide places it"),
        ("Final review (2026-10-29)", "all 260 rows verified up to that point were re-checked. Six were put back to AMBIGUOUS: four whose handle came from this guide's earlier Visitor Essentials record rather than a link on the venue's own page (the sources behind that record are aggregator or media pages), one whose recorded handle was a booking platform's own account, and one whose profile evidence could not tell the venue from other bars in the same block. Two rows confirmed from a page of the venue's own now cite that page instead of the Instagram URL."),
        ("Constraint on re-checking", "Instagram blocks unauthenticated profile reads, so a profile-based row cannot be re-opened and read again from here. Those rows were re-checked against what the research recorded at the time - whether a street address, a named station, a telephone number, an official marker or a branch name was written down - and rows without such a record were put back to AMBIGUOUS rather than left verified."),
        ("Per-venue search coverage", "all 338 listings that the website pass left unresolved have had the per-venue Instagram search run on them; all 111 rows that remained ambiguous have had a second per-venue pass, and all of those still ambiguous a third pass that read the sources recorded in the 557-row Maps audit list; every row carries either a result or the reason none could be established"),
        ("Why AMBIGUOUS exists", "an account is printed on a directory or media page, or several official accounts are linked and none is identifiably this venue's"),
        ("Why so many NOT_FOUND", "no official website is on record for the venue, and Instagram blocks unauthenticated profile reads, so no admissible evidence exists"),
        ("Checked on", next(iter(ledger.values()))["checked_at"]),
    ]
    for r in rows:
        s.append(list(r))
    s.column_dimensions["A"].width = 52
    s.column_dimensions["B"].width = 96
    for row in range(1, s.max_row + 1):
        s.cell(row=row, column=1).font = Font(bold=True)
        s.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT)
    print(f"written: {OUT}")
    print(f"rows: {ws.max_row - 1}  columns: {len(head)} ({len(MASTER_COLS)} original + {len(NEW_COLS)} new)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
